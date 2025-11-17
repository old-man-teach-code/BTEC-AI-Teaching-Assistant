import os
from pathlib import Path
from typing import List, Dict, Any, Optional
import openpyxl
import pandas as pd
from fastapi import UploadFile, HTTPException
from core.config import settings


async def read_excel_file(file: UploadFile) -> Dict[str, Any]:
    """
    Đọc nội dung file Excel
    
    Args:
        file: File Excel cần đọc
        
    Returns:
        Dict chứa dữ liệu từ file Excel
        
    Raises:
        HTTPException: Nếu không thể đọc file
    """
    try:
        # Đọc nội dung file
        contents = await file.read()
        
        # Xác định định dạng file
        file_extension = os.path.splitext(file.filename)[1].lower()
        
        if file_extension == '.xlsx':
            # Đọc file .xlsx với openpyxl
            workbook = openpyxl.load_workbook(filename=file.filename, data_only=True)
            result = {
                "filename": file.filename,
                "sheets": []
            }
            
            for sheet_name in workbook.sheetnames:
                sheet = workbook[sheet_name]
                data = []
                for row in sheet.iter_rows(values_only=True):
                    data.append(list(row))
                
                result["sheets"].append({
                    "name": sheet_name,
                    "data": data,
                    "rows": sheet.max_row,
                    "columns": sheet.max_column
                })
            
            return result
            
        elif file_extension == '.xls':
            # Đọc file .xls với pandas
            df_dict = pd.read_excel(contents, sheet_name=None, engine='xlrd')
            result = {
                "filename": file.filename,
                "sheets": []
            }
            
            for sheet_name, df in df_dict.items():
                result["sheets"].append({
                    "name": sheet_name,
                    "data": df.values.tolist(),
                    "columns": df.columns.tolist(),
                    "rows": len(df),
                    "columns_count": len(df.columns)
                })
            
            return result
        else:
            raise HTTPException(status_code=400, detail="Định dạng file không được hỗ trợ")
            
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Lỗi khi đọc file Excel: {str(e)}")
    finally:
        # Reset file pointer
        await file.seek(0)


async def validate_excel_structure(
    file: UploadFile,
    required_columns: Optional[List[str]] = None,
    sheet_name: Optional[str] = None
) -> Dict[str, Any]:
    """
    Kiểm tra cấu trúc file Excel
    
    Args:
        file: File Excel cần kiểm tra
        required_columns: Danh sách các cột bắt buộc (nếu có)
        sheet_name: Tên sheet cần kiểm tra (nếu không chỉ định sẽ lấy sheet đầu tiên)
        
    Returns:
        Dict chứa kết quả kiểm tra
    """
    try:
        excel_data = await read_excel_file(file)
        
        # Lấy sheet cần kiểm tra
        if sheet_name:
            sheet_data = next((s for s in excel_data["sheets"] if s["name"] == sheet_name), None)
            if not sheet_data:
                return {
                    "valid": False,
                    "message": f"Không tìm thấy sheet '{sheet_name}'"
                }
        else:
            sheet_data = excel_data["sheets"][0] if excel_data["sheets"] else None
            
        if not sheet_data:
            return {
                "valid": False,
                "message": "File Excel không có sheet nào"
            }
        
        # Kiểm tra các cột bắt buộc
        if required_columns:
            # Lấy header (dòng đầu tiên)
            if sheet_data["data"]:
                headers = [str(col).lower() if col else "" for col in sheet_data["data"][0]]
                missing_columns = []
                
                for required_col in required_columns:
                    if required_col.lower() not in headers:
                        missing_columns.append(required_col)
                
                if missing_columns:
                    return {
                        "valid": False,
                        "message": f"Thiếu các cột: {', '.join(missing_columns)}"
                    }
        
        return {
            "valid": True,
            "message": "File Excel hợp lệ",
            "sheet_name": sheet_data["name"],
            "rows": sheet_data["rows"],
            "columns": sheet_data.get("columns_count", sheet_data.get("columns", 0))
        }
        
    except Exception as e:
        return {
            "valid": False,
            "message": f"Lỗi khi kiểm tra file Excel: {str(e)}"
        }


async def convert_excel_to_json(file: UploadFile, sheet_name: Optional[str] = None) -> List[Dict[str, Any]]:
    """
    Chuyển đổi file Excel sang định dạng JSON
    
    Args:
        file: File Excel cần chuyển đổi
        sheet_name: Tên sheet cần chuyển đổi (nếu không chỉ định sẽ lấy sheet đầu tiên)
        
    Returns:
        List các dict đại diện cho từng dòng trong Excel
    """
    try:
        contents = await file.read()
        await file.seek(0)
        
        # Đọc file Excel với pandas
        if sheet_name:
            df = pd.read_excel(contents, sheet_name=sheet_name)
        else:
            df = pd.read_excel(contents, sheet_name=0)
        
        # Chuyển đổi sang JSON (list of dicts)
        result = df.to_dict(orient='records')
        
        # Làm sạch các giá trị NaN
        for record in result:
            for key, value in record.items():
                if pd.isna(value):
                    record[key] = None
        
        return result
        
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Lỗi khi chuyển đổi file Excel: {str(e)}")


def create_excel_file(data: List[Dict[str, Any]], filename: str, sheet_name: str = "Sheet1") -> str:
    """
    Tạo file Excel từ dữ liệu
    
    Args:
        data: Dữ liệu cần ghi vào Excel (list of dicts)
        filename: Tên file Excel
        sheet_name: Tên sheet
        
    Returns:
        str: Đường dẫn đến file Excel đã tạo
    """
    try:
        # Tạo DataFrame từ data
        df = pd.DataFrame(data)
        
        # Tạo đường dẫn file
        upload_dir = Path(settings.UPLOAD_DIR)
        upload_dir.mkdir(parents=True, exist_ok=True)
        
        file_path = upload_dir / filename
        
        # Ghi vào file Excel
        with pd.ExcelWriter(file_path, engine='openpyxl') as writer:
            df.to_excel(writer, sheet_name=sheet_name, index=False)
        
        return str(file_path)
        
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Lỗi khi tạo file Excel: {str(e)}")
